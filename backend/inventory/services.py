from datetime import datetime
from django.db import transaction
from rest_framework.exceptions import ValidationError
from django.db.models import Max
from openpyxl import load_workbook

import csv
import io
import os

from .models import (Car, CarImage, CarExpense,)


class InventoryService:
    
    STATUS_TRANSITIONS = {
        "available": {
            "available",
            "upcoming",
            "reserved",
            "in_service",
            "in_house",
        },
        "upcoming": {
            "upcoming",
            "available",
            "in_service",
            "in_house",
        },
        "reserved": {
            "reserved",
            "available",
            "booked",
        },
        "booked": {
            "booked",
            "reserved",
            "sold",
        },
        "sold": {
            "sold",
        },
        "in_service": {
            "in_service",
            "available",
            "in_house",
        },
        "in_house": {
            "in_house",
            "available",
            "in_service",
        },
    }

    @staticmethod
    def validate_status_transition(car, new_status):
        current_status = car.status

        if current_status == new_status:
            return

        allowed_statuses = InventoryService.STATUS_TRANSITIONS.get(
            current_status,
            set(),
        )

        if new_status not in allowed_statuses:
            raise ValidationError(
                {
                    "status": (
                        f"Invalid vehicle status transition: "
                        f"{current_status} → {new_status}."
                    )
                }
            )

    @staticmethod
    def validate_service_location(status, service_location):
        if status == "in_service" and not service_location:
            raise ValidationError(
                {
                    "service_location": (
                        "Service location is required when "
                        "vehicle status is in service."
                    )
                }
            )

    @staticmethod
    def generate_stock_id():

        latest = (
            Car.objects
            .filter(stock_id__startswith="PR-L")
            .aggregate(max_stock=Max("stock_id"))
        )["max_stock"]

        if latest:
            sequence = int(latest.replace("PR-L", "")) + 1
        else:
            sequence = 1

        return f"PR-L{sequence:04d}"

    @staticmethod
    def _save_car_images(
        *,
        car,
        images,
    ):
        """
        Save uploaded gallery images for a vehicle.
        """

        if not images:
            return

        for image in images:
            CarImage.objects.create(
                car=car,
                image=image,
                is_cover=False,
            )

    @staticmethod
    def validate_bulk_vehicle_identifiers(
        vehicles,
        row_numbers=None,
    ):
        """
        Validate chassis_number and engine_number for a bulk import.

        Checks:

        1. Duplicate chassis numbers inside the import file.
        2. Duplicate engine numbers inside the import file.
        3. Chassis numbers already existing in the database.
        4. Engine numbers already existing in the database.

        Returns:
            list[dict]: Row-level errors.
        """

        errors = []

        # ---------------------------------------------------------
        # Generate row numbers if not explicitly supplied
        # ---------------------------------------------------------

        if row_numbers is None:
            row_numbers = list(
                range(
                    2,
                    len(vehicles) + 2,
                )
            )

        if len(row_numbers) != len(vehicles):
            raise ValueError(
                "row_numbers length must match vehicles length."
            )

        chassis_seen = {}
        engine_seen = {}

        # =========================================================
        # NORMALIZE + DETECT DUPLICATES INSIDE IMPORT
        # =========================================================

        for index, vehicle in enumerate(vehicles):

            row_number = row_numbers[index]

            # -----------------------------------------------------
            # CHASSIS
            # -----------------------------------------------------

            chassis = vehicle.get(
                "chassis_number"
            )

            if chassis:

                chassis = (
                    str(chassis)
                    .strip()
                    .upper()
                )

                vehicle["chassis_number"] = chassis

                if chassis in chassis_seen:

                    errors.append(
                        {
                            "row": row_number,
                            "field": "chassis_number",
                            "value": chassis,
                            "message": (
                                f"Chassis number {chassis} "
                                f"is duplicated in the import file. "
                                f"It was already found on row "
                                f"{chassis_seen[chassis]}."
                            ),
                        }
                    )

                else:

                    # Store ACTUAL Excel row number
                    chassis_seen[chassis] = row_number

            # -----------------------------------------------------
            # ENGINE
            # -----------------------------------------------------

            engine = vehicle.get(
                "engine_number"
            )

            if engine:

                engine = (
                    str(engine)
                    .strip()
                    .upper()
                )

                vehicle["engine_number"] = engine

                if engine in engine_seen:

                    errors.append(
                        {
                            "row": row_number,
                            "field": "engine_number",
                            "value": engine,
                            "message": (
                                f"Engine number {engine} "
                                f"is duplicated in the import file. "
                                f"It was already found on row "
                                f"{engine_seen[engine]}."
                            ),
                        }
                    )

                else:

                    # Store ACTUAL Excel row number
                    engine_seen[engine] = row_number

        # =========================================================
        # DATABASE CHECK
        # =========================================================

        chassis_numbers = list(
            chassis_seen.keys()
        )

        engine_numbers = list(
            engine_seen.keys()
        )

        # ---------------------------------------------------------
        # Existing chassis numbers
        # ---------------------------------------------------------

        existing_chassis = {
            value.upper()
            for value in (
                Car.objects
                .filter(
                    chassis_number__in=chassis_numbers
                )
                .values_list(
                    "chassis_number",
                    flat=True,
                )
            )
            if value
        }

        # ---------------------------------------------------------
        # Existing engine numbers
        # ---------------------------------------------------------

        existing_engines = {
            value.upper()
            for value in (
                Car.objects
                .filter(
                    engine_number__in=engine_numbers
                )
                .values_list(
                    "engine_number",
                    flat=True,
                )
            )
            if value
        }

        # =========================================================
        # REPORT EXISTING CHASSIS
        # =========================================================

        for index, vehicle in enumerate(vehicles):

            row_number = row_numbers[index]

            chassis = vehicle.get(
                "chassis_number"
            )

            if not chassis:
                continue

            chassis = (
                str(chassis)
                .strip()
                .upper()
            )

            if chassis in existing_chassis:

                existing_car = (
                    Car.objects
                    .filter(
                        chassis_number__iexact=chassis
                    )
                    .only(
                        "stock_id",
                        "chassis_number",
                    )
                    .first()
                )

                stock_id = (
                    existing_car.stock_id
                    if existing_car
                    else "unknown"
                )

                errors.append(
                    {
                        "row": row_number,
                        "field": "chassis_number",
                        "value": chassis,
                        "message": (
                            f"Chassis number {chassis} "
                            f"already exists for "
                            f"stock {stock_id}."
                        ),
                    }
                )

        # =========================================================
        # REPORT EXISTING ENGINE
        # =========================================================

        for index, vehicle in enumerate(vehicles):

            row_number = row_numbers[index]

            engine = vehicle.get(
                "engine_number"
            )

            if not engine:
                continue

            engine = (
                str(engine)
                .strip()
                .upper()
            )

            if engine in existing_engines:

                existing_car = (
                    Car.objects
                    .filter(
                        engine_number__iexact=engine
                    )
                    .only(
                        "stock_id",
                        "engine_number",
                    )
                    .first()
                )

                stock_id = (
                    existing_car.stock_id
                    if existing_car
                    else "unknown"
                )

                errors.append(
                    {
                        "row": row_number,
                        "field": "engine_number",
                        "value": engine,
                        "message": (
                            f"Engine number {engine} "
                            f"already exists for "
                            f"stock {stock_id}."
                        ),
                    }
                )

        return errors


    @staticmethod
    @transaction.atomic
    def bulk_create_cars(
        vehicles,
    ):
        """
        Create all vehicles atomically.

        If any vehicle fails during creation,
        the entire transaction is rolled back.
        """

        created_cars = []

        for vehicle in vehicles:

            vehicle_data = dict(
                vehicle
            )

            # Remove fields that are not Car model fields
            vehicle_data.pop(
                "images",
                None,
            )

            # Generate stock ID for each vehicle
            vehicle_data["stock_id"] = (
                InventoryService
                .generate_stock_id()
            )

            car = Car.objects.create(
                **vehicle_data
            )

            created_cars.append(
                car
            )

        return created_cars


    @staticmethod
    def import_cars(
        vehicles,
        start_row=2,
    ):
        """
        Full production bulk-import workflow.

        Flow:

            Input
              ↓
            Normalize
              ↓
            Validate identifiers
              ↓
            If errors → reject entire batch
              ↓
            Atomic creation
        """

        # =====================================================
        # COPY INPUT
        # =====================================================

        normalized_vehicles = []

        for vehicle in vehicles:

            vehicle_data = dict(
                vehicle
            )

            vehicle_data = (
                InventoryService
                .normalize_vehicle_data(
                    vehicle_data
                )
            )

            normalized_vehicles.append(
                vehicle_data
            )

        # =====================================================
        # IDENTIFIER VALIDATION
        # =====================================================

            errors = (
            InventoryService
            .validate_bulk_vehicle_identifiers(
                normalized_vehicles,
                row_numbers=list(
                    range(
                        start_row,
                        start_row + len(normalized_vehicles),
                    )
                ),
            )
        )

        # =====================================================
        # REJECT ENTIRE BATCH
        # =====================================================

        if errors:

            raise ValidationError(
                {
                    "bulk_import": errors
                }
            )

        # =====================================================
        # CREATE ATOMICALLY
        # =====================================================

        return (
            InventoryService
            .bulk_create_cars(
                normalized_vehicles
            )
        )
    
    @staticmethod
    def normalize_vehicle_data(validated_data):
        """
        Normalize vehicle data before saving.

        Missing text fields are stored as empty strings.
        """

        validated_data = dict(
            validated_data
        )

        text_fields = [
            "make",
            "model",
            "variant",
            "colour",
            "chassis_number",
            "engine_number",
            "source_specify",
            "supplier",
        ]

        for field in text_fields:

            value = validated_data.get(
                field
            )

            if value is None:
                validated_data[field] = ""
                continue

            value = str(value).strip()

            if field in [
                "chassis_number",
                "engine_number",
            ]:
                value = value.upper()

            elif field in [
                "make",
                "model",
                "variant",
                "colour",
                "supplier",
            ]:
                value = value.title()

            validated_data[field] = value

        return validated_data

    @staticmethod
    def create_car(validated_data):
        images = validated_data.pop(
            "images",
            [],
        )

        validated_data = InventoryService.normalize_vehicle_data(
            validated_data,
        )

        status = validated_data.get(
            "status",
            "available",
        )

        service_location = validated_data.get(
            "service_location",
            "",
        )

        InventoryService.validate_service_location(
            status,
            service_location,
        )

        validated_data["stock_id"] = (
            InventoryService.generate_stock_id()
        )

        car = Car.objects.create(
            **validated_data
        )

        InventoryService._save_car_images(
            car=car,
            images=images,
        )

        return car
    
    @staticmethod
    def update_car(
        car,
        validated_data,
    ):

        remove_certificate = validated_data.pop(
        "remove_certificate",
        False,
    )

        images = validated_data.pop(
            "images",
            [],
        )

        text_fields = [
            "make",
            "model",
            "variant",
            "colour",
            "chassis_number",
            "engine_number",
            "source_specify",
            "supplier",
        ]

        for field in text_fields:
            if field not in validated_data:
                continue

            value = validated_data[field]

            if value is None:
                validated_data[field] = ""
                continue

            value = str(value).strip()

            if field in [
                "chassis_number",
                "engine_number",
            ]:
                value = value.upper()

            elif field in [
                "make",
                "model",
                "variant",
                "colour",
                "supplier",
            ]:
                value = value.title()

            validated_data[field] = value

        new_status = validated_data.get(
            "status",
            car.status,
        )

        new_service_location = validated_data.get(
            "service_location",
            car.service_location,
        )

        InventoryService.validate_status_transition(
            car,
            new_status,
        )

        InventoryService.validate_service_location(
            new_status,
            new_service_location,
        )


        if remove_certificate:

            if car.possession_certificate:

                car.possession_certificate.delete(
                    save=False,
                )

            car.possession_certificate = None

        for field, value in validated_data.items():
            setattr(
                car,
                field,
                value,
            )

        car.save()

        InventoryService._save_car_images(
            car=car,
            images=images,
        )

        return car


    @staticmethod
    def set_cover_image(
        *,
        image,
    ):
        """
        Set one image as the cover image for a vehicle.
        """

        CarImage.objects.filter(
            car=image.car,
            is_cover=True,
        ).update(
            is_cover=False,
        )

        image.is_cover = True

        image.save(
            update_fields=[
                "is_cover",
            ]
        )

        return image

    @staticmethod
    def update_expense(
        expense,
        validated_data,
    ):

        for field, value in validated_data.items():
            setattr(
                expense,
                field,
                value,
            )

        expense.save()

        return expense


    @staticmethod
    def delete_expense(
        expense,
    ):

        expense.delete()

    @staticmethod
    def delete_car(
        car,
    ):

        # Delete certificate file
        if car.possession_certificate:
            car.possession_certificate.delete(
                save=False,
            )

        # Delete image files
        for image in car.images.all():
            image.image.delete(
                save=False,
            )

        car.delete()

    @staticmethod
    def delete_image(image):

        car = image.car

        was_cover = image.is_cover

        image.image.delete(save=False)

        image.delete()

        if was_cover:

            new_cover = (
                CarImage.objects
                .filter(car=car)
                .first()
            )

            if new_cover:

                new_cover.is_cover = True

                new_cover.save(
                    update_fields=[
                        "is_cover",
                    ]
                )


    @staticmethod
    @transaction.atomic
    def reorder_images(
        car,
        images,
        image_order,
    ):
        """
        Reorders images for a vehicle.

        Parameters
        ----------
        car : Car
            Vehicle whose images are being reordered.

        images : QuerySet[CarImage]
            Images belonging to this vehicle.

        image_order : list[int]
            Ordered list of image IDs.
        """

        images_map = {
            image.id: image
            for image in images
        }

        # Validate count
        if len(image_order) != len(images_map):
            raise ValidationError(
                {
                    "image_order": [
                        "Image list is incomplete."
                    ]
                }
            )

        # Validate duplicates
        if len(image_order) != len(set(image_order)):
            raise ValidationError(
                {
                    "image_order": [
                        "Duplicate image IDs detected."
                    ]
                }
            )

        # Validate ownership
        for image_id in image_order:

            if image_id not in images_map:

                raise ValidationError(
                {
                    "image_order": [
                        "One or more images do not belong to this vehicle."
                    ]
                }
            )

        # Update display order
        for order, image_id in enumerate(
            image_order,
            start=1,
        ):

            image = images_map[image_id]

            image.display_order = order

            image.save(
                update_fields=[
                    "display_order",
                ]
            )

        return images

    
    @staticmethod
    @transaction.atomic
    def bulk_delete_images(images):

        deleted_count = 0

        for image in images:

            InventoryService.delete_image(
                image,
            )

            deleted_count += 1

        return deleted_count

    @staticmethod
    @transaction.atomic
    def bulk_delete_vehicles(
        cars,
    ):

        deleted = 0

        for car in cars:

            InventoryService.delete_car(
                car,
            )

            deleted += 1

        return deleted
    
    @staticmethod
    @transaction.atomic
    def delete_all_cars(cars):
        """
        Delete the complete inventory stock atomically.

        The database deletion is completed first. Physical files are
        removed only after the transaction successfully commits.

        If any vehicle is protected by an existing business record,
        the entire database operation is rolled back.
        """

        cars = list(
            cars.select_for_update()
        )

        if not cars:
            return 0

        files_to_delete = []

        for car in cars:

            if car.possession_certificate:
                files_to_delete.append(
                    car.possession_certificate.name
                )

            for image in car.images.all():

                if image.image:
                    files_to_delete.append(
                        image.image.name
                    )

            for document in car.vehicle_documents.all():

                if document.file:
                    files_to_delete.append(
                        document.file.name
                    )

        deleted_count = len(cars)

        # Database deletion happens before physical file deletion.
        # PROTECT relationships will raise an exception here and
        # transaction.atomic will roll back the entire operation.
        Car.objects.filter(
            id__in=[car.id for car in cars]
        ).delete()

        def delete_files():

            from django.core.files.storage import default_storage

            for file_name in files_to_delete:

                if file_name and default_storage.exists(file_name):
                    default_storage.delete(file_name)

        transaction.on_commit(
            delete_files
        )

        return deleted_count
    @staticmethod
    def parse_vehicle_import_file(file):
        """
        Parse an XLSX or CSV vehicle import file.

        Returns:
            list[dict]: Vehicle rows.

        Each row has:
            {
                "row_number": 2,
                "data": {...}
            }
        """

        filename = (
            getattr(file, "name", "") or ""
        ).lower()

        # =====================================================
        # CSV
        # =====================================================

        if filename.endswith(".csv"):

            try:
                raw_data = file.read()

                text = raw_data.decode(
                    "utf-8-sig"
                )

                reader = csv.DictReader(
                    io.StringIO(text)
                )

            except UnicodeDecodeError:

                raise ValidationError(
                    {
                        "file": [
                            "Unable to read CSV file. "
                            "CSV must be UTF-8 encoded."
                        ]
                    }
                )

            except Exception as exc:

                raise ValidationError(
                    {
                        "file": [
                            f"Unable to read CSV file: {str(exc)}"
                        ]
                    }
                )

            # -------------------------------------------------
            # Empty CSV / missing headers
            # -------------------------------------------------

            if not reader.fieldnames:

                raise ValidationError(
                    {
                        "file": [
                            "The CSV file is empty or "
                            "contains no header row."
                        ]
                    }
                )

            headers = [
                str(value).strip()
                if value is not None
                else ""
                for value in reader.fieldnames
            ]

            reader.fieldnames = headers

            # -------------------------------------------------
            # Keep the same columns as XLSX
            # -------------------------------------------------

            required_columns = [
                "make",
                "model",
                "variant",
                "colour",
                "year",
                "purchase_cost",
                "asking_price",
                "least_selling_price",
                "status",
                "mileage",
                "source",
                "source_specify",
                "supplier",
                "highlight_public",
                "chassis_number",
                "engine_number",
            ]

            missing_columns = [
                field
                for field in required_columns
                if field not in headers
            ]

            if missing_columns:

                raise ValidationError(
                    {
                        "file": [
                            "Missing required columns: "
                            + ", ".join(
                                missing_columns
                            )
                        ]
                    }
                )

            vehicles = []

            # -------------------------------------------------
            # Read CSV rows
            # -------------------------------------------------

            for row_number, row in enumerate(
                reader,
                start=2,
            ):

                # Ignore completely empty rows
                if not any(
                    value is not None
                    and str(value).strip() != ""
                    for value in row.values()
                ):
                    continue

                vehicle = {}

                for field in required_columns:

                    value = row.get(
                        field
                    )

                    if isinstance(
                        value,
                        str,
                    ):

                        value = value.strip()

                        if value == "":
                            value = None

                    vehicle[field] = value

                vehicles.append(
                    {
                        "row_number": row_number,
                        "data": vehicle,
                    }
                )

            if not vehicles:

                raise ValidationError(
                    {
                        "file": [
                            "The CSV file contains "
                            "no vehicle rows."
                        ]
                    }
                )

            return vehicles

        # =====================================================
        # XLSX
        # =====================================================

        if filename.endswith(".xlsx"):

            try:

                workbook = load_workbook(
                    file,
                    read_only=True,
                    data_only=True,
                )

            except Exception as exc:

                raise ValidationError(
                    {
                        "file": [
                            f"Unable to read Excel file: {str(exc)}"
                        ]
                    }
                )

            try:

                worksheet = workbook.active

                rows = list(
                    worksheet.iter_rows(
                        values_only=True,
                    )
                )

                if not rows:

                    raise ValidationError(
                        {
                            "file": [
                                "The Excel file is empty."
                            ]
                        }
                    )

                headers = [
                    str(value).strip()
                    if value is not None
                    else ""
                    for value in rows[0]
                ]

                required_columns = [
                    "make",
                    "model",
                    "variant",
                    "colour",
                    "year",
                    "purchase_cost",
                    "asking_price",
                    "least_selling_price",
                    "status",
                    "mileage",
                    "source",
                    "source_specify",
                    "supplier",
                    "highlight_public",
                    "chassis_number",
                    "engine_number",
                ]

                missing_columns = [
                    field
                    for field in required_columns
                    if field not in headers
                ]

                if missing_columns:

                    raise ValidationError(
                        {
                            "file": [
                                "Missing required columns: "
                                + ", ".join(
                                    missing_columns
                                )
                            ]
                        }
                    )

                header_map = {
                    header: index
                    for index, header in enumerate(
                        headers
                    )
                }

                vehicles = []

                for row_number, row in enumerate(
                    rows[1:],
                    start=2,
                ):

                    # Ignore completely empty rows
                    if not any(
                        value is not None
                        and str(value).strip() != ""
                        for value in row
                    ):
                        continue

                    vehicle = {}

                    for field in required_columns:

                        column_index = (
                            header_map[field]
                        )

                        value = (
                            row[column_index]
                            if column_index < len(row)
                            else None
                        )

                        if isinstance(
                            value,
                            str,
                        ):

                            value = value.strip()

                            if value == "":
                                value = None

                        vehicle[field] = value

                    vehicles.append(
                        {
                            "row_number": row_number,
                            "data": vehicle,
                        }
                    )

                if not vehicles:

                    raise ValidationError(
                        {
                            "file": [
                                "The Excel file contains "
                                "no vehicle rows."
                            ]
                        }
                    )

                return vehicles

            finally:

                workbook.close()

        # =====================================================
        # Unsupported file type
        # =====================================================

        raise ValidationError(
            {
                "file": [
                    "Unsupported file format. "
                    "Please upload an .xlsx or .csv file."
                ]
            }
        )


    @staticmethod
    def validate_vehicle_import_rows(rows):
        """
        Validate all vehicle rows before creation.

        Bulk import rules:

        Required:
            year
            make
            model
            variant
            colour

        Optional:
            chassis_number
            engine_number
            purchase_cost
            asking_price
            least_selling_price
            mileage
            source_specify
            supplier
            etc.

        Invalid rows are reported individually.
        Valid rows remain eligible for creation.
        """

        normalized_vehicles = []
        errors = []

        # -----------------------------------------------------
        # Normalize rows
        # -----------------------------------------------------

        for item in rows:

            row_number = item["row_number"]

            vehicle = dict(
                item["data"]
            )

            vehicle = (
                InventoryService
                .normalize_vehicle_data(
                    vehicle
                )
            )

            # IMPORTANT:
            # Preserve the actual Excel row number.
            vehicle["_row_number"] = row_number

            normalized_vehicles.append(
                vehicle
            )

        # -----------------------------------------------------
        # Identifier validation
        # -----------------------------------------------------

        row_numbers = [
            vehicle["_row_number"]
            for vehicle in normalized_vehicles
        ]

        identifier_errors = (
            InventoryService
            .validate_bulk_vehicle_identifiers(
                normalized_vehicles,
                row_numbers=row_numbers,
            )
        )

        errors.extend(
            identifier_errors
        )

        # -----------------------------------------------------
        # Validate every row through the bulk serializer
        # -----------------------------------------------------

        from .serializers import (
            BulkVehicleRowSerializer,
        )

        for vehicle in normalized_vehicles:

            row_number = vehicle["_row_number"]

            serializer_data = dict(
                vehicle
            )

            # Internal field.
            serializer_data.pop(
                "_row_number",
                None,
            )

            serializer = BulkVehicleRowSerializer(
                data=serializer_data
            )

            if not serializer.is_valid():

                for field, messages in (
                    serializer.errors.items()
                ):

                    if not isinstance(
                        messages,
                        list,
                    ):
                        messages = [messages]

                    for message in messages:

                        errors.append(
                            {
                                "row": row_number,
                                "field": field,
                                "value": vehicle.get(field),
                                "message": str(message),
                            }
                        )

        return (
            normalized_vehicles,
            errors,
        )
    @staticmethod
    def import_vehicle_rows(rows):
        """
        Import vehicles independently.

        Valid rows are created.
        Invalid rows are skipped and reported.

        One invalid vehicle does NOT prevent
        other valid vehicles from being imported.
        """

        (
            normalized_vehicles,
            errors,
        ) = (
            InventoryService
            .validate_vehicle_import_rows(
                rows
            )
        )

        # -----------------------------------------------------
        # Group validation errors by Excel row
        # -----------------------------------------------------

        errors_by_row = {}

        for error in errors:

            row_number = str(
                error.get("row")
            )

            errors_by_row.setdefault(
                row_number,
                []
            ).append(
                error
            )

        created_cars = []
        skipped_rows = []

        # -----------------------------------------------------
        # Process every row independently
        # -----------------------------------------------------

        for vehicle_data in normalized_vehicles:

            vehicle_data = dict(
                vehicle_data
            )

            # -------------------------------------------------
            # Get actual Excel row number
            # -------------------------------------------------

            row_number = str(
                vehicle_data.pop(
                    "_row_number"
                )
            )

            # -------------------------------------------------
            # Skip rows that already failed validation
            # -------------------------------------------------

            if row_number in errors_by_row:

                skipped_rows.append(
                    row_number
                )

                continue

            # -------------------------------------------------
            # REQUIRED VEHICLE IDENTITY SAFETY CHECK
            #
            # Defense-in-depth.
            #
            # Do NOT raise ValidationError here.
            # This row must be skipped while other rows
            # continue importing.
            # -------------------------------------------------

            required_fields = {
                "year": "Year is required.",
                "make": "Make is required.",
                "model": "Model is required.",
                "variant": "Variant is required.",
                "colour": "Colour is required.",
            }

            missing_field = None

            for field, message in (
                required_fields.items()
            ):

                value = vehicle_data.get(
                    field
                )

                if value is None:

                    missing_field = {
                        "field": field,
                        "value": None,
                        "message": message,
                    }

                    break

                if (
                    isinstance(value, str)
                    and not value.strip()
                ):

                    missing_field = {
                        "field": field,
                        "value": value,
                        "message": message,
                    }

                    break

            # -------------------------------------------------
            # Reject only this row
            # -------------------------------------------------

            if missing_field:

                error = {
                    "row": row_number,
                    "field": missing_field["field"],
                    "value": missing_field["value"],
                    "message": missing_field["message"],
                }

                errors_by_row.setdefault(
                    row_number,
                    []
                ).append(
                    error
                )

                skipped_rows.append(
                    row_number
                )

                continue

            # -------------------------------------------------
            # Optional text fields
            # -------------------------------------------------

            optional_text_fields = [
                "chassis_number",
                "engine_number",
                "source_specify",
                "supplier",
            ]

            for field in optional_text_fields:

                if vehicle_data.get(field) is None:

                    vehicle_data[field] = ""

            # -------------------------------------------------
            # Never pass import-only fields to Car
            # -------------------------------------------------

            vehicle_data.pop(
                "images",
                None,
            )

            vehicle_data.pop(
                "_row_number",
                None,
            )

            # -------------------------------------------------
            # Generate stock ID
            # -------------------------------------------------

            vehicle_data["stock_id"] = (
                InventoryService
                .generate_stock_id()
            )

            # -------------------------------------------------
            # Create ONLY this vehicle atomically
            # -------------------------------------------------

            try:

                with transaction.atomic():

                    car = Car.objects.create(
                        **vehicle_data
                    )

                created_cars.append(
                    car
                )

            except Exception as exc:

                skipped_rows.append(
                    row_number
                )

                errors_by_row.setdefault(
                    row_number,
                    []
                ).append(
                    {
                        "row": row_number,
                        "field": "database",
                        "value": None,
                        "message": str(exc),
                    }
                )

        # -----------------------------------------------------
        # Flatten errors
        # -----------------------------------------------------

        all_errors = [
            error
            for row_errors in (
                errors_by_row.values()
            )
            for error in row_errors
        ]

        return {
            "created_cars": created_cars,
            "created_count": len(
                created_cars
            ),
            "skipped_count": len(
                skipped_rows
            ),
            "skipped_rows": skipped_rows,
            "errors": all_errors,
        }
        
    @staticmethod
    @transaction.atomic
    def update_car_status(
        *,
        car,
        new_status,
    ):
        InventoryService.validate_status_transition(
            car,
            new_status,
        )

        car.status = new_status

        car.save(
            update_fields=[
                "status",
                "updated_at",
            ]
        )

        return car
    
    @staticmethod
    @transaction.atomic
    def prepare_car_for_reservation(
        *,
        car,
    ):
        """
        Prepare a vehicle for reservation using the existing
        Inventory status transition rules.

        In-house vehicles must first become Available.
        Available vehicles then become Reserved.

        Existing Reserved vehicles are left unchanged.
        """

        if car.status == Car.Status.IN_HOUSE:
            InventoryService.update_car_status(
                car=car,
                new_status=Car.Status.AVAILABLE,
            )

        if car.status == Car.Status.AVAILABLE:
            InventoryService.update_car_status(
                car=car,
                new_status=Car.Status.RESERVED,
            )

        return car
    
    @staticmethod
    @transaction.atomic
    def prepare_car_for_booking(
        *,
        car,
    ):
        """
        Prepare a vehicle for booking using the existing
        Inventory status transition rules.

        Valid progression:
            In House → Available → Reserved → Booked

        Existing Booked vehicles are left unchanged.
        """

        if car.status == Car.Status.IN_HOUSE:
            InventoryService.update_car_status(
                car=car,
                new_status=Car.Status.AVAILABLE,
            )

        if car.status == Car.Status.AVAILABLE:
            InventoryService.update_car_status(
                car=car,
                new_status=Car.Status.RESERVED,
            )

        if car.status == Car.Status.RESERVED:
            InventoryService.update_car_status(
                car=car,
                new_status=Car.Status.BOOKED,
            )

        return car